#!/usr/bin/env python3
"""生成酒店知识采集模板与脱敏示例 Excel。

用法（仓库根目录）：
    python3 tools/joctv-hotel-kb-skill/scripts/make_excel_template.py

输出：
    tools/joctv-hotel-kb-skill/templates/酒店知识采集模板.xlsx
    tools/joctv-hotel-kb-skill/examples/脱敏示例.xlsx
需要 openpyxl（pip install openpyxl）。
"""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    print("缺少 openpyxl，请先: pip install openpyxl")
    sys.exit(2)

SKILL_ROOT = Path(__file__).resolve().parent.parent
TEMPLATE_PATH = SKILL_ROOT / "templates" / "酒店知识采集模板.xlsx"
EXAMPLE_PATH = SKILL_ROOT / "examples" / "脱敏示例.xlsx"

CATEGORIES = ["衣", "食", "住", "行", "游", "购", "玩", "服务", "其他"]

CATEGORY_TABLE = [
    ("衣", "Clothing & Laundry", "洗衣、熨烫、衣物寄存等"),
    ("食", "Food & Drink", "餐厅、酒吧、早餐、下午茶、客房送餐"),
    ("住", "Rooms & Stay", "客房、套房、房型、入住退房"),
    ("行", "Transport", "地铁、机场、接驳、停车、路线"),
    ("游", "Attractions", "周边景点、游览建议"),
    ("购", "Shopping", "商圈、购物中心、免税"),
    ("玩", "Wellness & Fun", "健身房、泳池、水疗、儿童乐园"),
    ("服务", "Services", "前台电话、政策、会议婚宴、宠物、叫醒"),
    ("其他", "Others", "酒店概览、历史荣誉等不适合归入以上分类的"),
]

# (字段名, 必填, 填写说明)
INSTRUCTIONS = [
    ("总则", "", "一行 = 一个知识主题。请用酒店人员平时的叫法填写，不用写英文的地方可以留空。"),
    ("总则", "", "带 * 的列必填；“中文标准回答”与“英文标准回答”至少填一个，另一种语言可由公司同事帮忙补齐。"),
    ("总则", "", "回答可以填多条，在单元格内换行分隔；公司工作站会把每条回答扩写成同一组事实、至少 3 种自然改写（互补事实会归入备注，不会凑成不同答案）。"),
    ("总则", "", "只填酒店确认过的信息：时间、价格、地址、电话、政策不要估、不要抄别的酒店。没有就留空。"),
    ("编号", "可不填", "留空即可，导入系统时自动生成。"),
    ("分类 *", "必填", "从九类里选一个：衣、食、住、行、游、购、玩、服务、其他（见“分类表”Sheet）。"),
    ("中文主题 *", "必填", "这个知识叫什么，如：健身房、唐阁、机场接送。"),
    ("英文主题", "尽量填", "对应的英文名，如：Fitness Center。不确定可留空。"),
    ("中文命中词", "尽量填", "宾客可能怎么问，多个词用分号（;）隔开，如：健身房;健身中心;gym在哪。"),
    ("英文命中词", "尽量填", "英文叫法，分号隔开，如：gym;fitness center。"),
    ("中文标准回答 *", "必填（与英文至少一个）", "标准答案，写宾客听得懂的话，可换行填多条。"),
    ("英文标准回答", "可空", "英文标准答案，可换行填多条；没有可留空，不要把中文贴进来。"),
    ("时间（中文）/ 时间（英文）", "可空", "营业时间、服务时间，如：每天 6:30-23:00。"),
    ("地点（中文）/ 地点（英文）", "可空", "楼层、区域或地址，如：2楼；B1层。"),
    ("怎么去（中文）/ 怎么去（英文）", "可空", "路线说明，如：出电梯左转到底。"),
    ("电话", "可空", "前台、餐厅或服务电话，如：86 (21) 1234 5678。"),
    ("其他备注（中文）/ 其他备注（英文）", "可空", "预约方式、人数限制、着装要求、教练、特色等其他事实。"),
    ("示例行", "", "知识条目Sheet里的灰色示例行仅供参照，正式填写前请删除。"),
]

ENTRY_COLUMNS = [
    "编号\n（可不填）",
    "分类 *",
    "中文主题 *",
    "英文主题",
    "中文命中词\n（分号分隔）",
    "英文命中词\n（分号分隔）",
    "中文标准回答 *\n（可换行填多条）",
    "英文标准回答\n（可换行填多条）",
    "时间（中文）",
    "时间（英文）",
    "地点（中文）",
    "地点（英文）",
    "怎么去（中文）",
    "怎么去（英文）",
    "电话",
    "其他备注（中文）",
    "其他备注（英文）",
]

TEMPLATE_EXAMPLE_ROWS = [
    [
        "",
        "玩",
        "健身房",
        "Fitness Center",
        "健身房;健身中心",
        "gym;fitness center",
        "健身房在2楼，每天6:00到23:00开放，免费使用。",
        "The gym is on the 2nd floor, open daily from 6:00 to 23:00, free of charge.",
        "每天 6:00-23:00",
        "Daily 6:00-23:00",
        "2楼",
        "2nd floor",
        "大堂电梯上2楼即到",
        "Take the lobby elevator to the 2nd floor",
        "86 (21) 1234 5678",
        "有跑步机和自由重量区；12岁以下需成人陪同。（此行为示例，请删除）",
        "Treadmills and free weights; under-12s need an adult. (Example row - please delete)",
    ],
    [
        "",
        "食",
        "早餐",
        "Breakfast",
        "早餐;几点吃早饭;自助餐",
        "breakfast;buffet",
        "自助早餐在1楼餐厅，周一到周五 7:00-10:30，周末 7:00-11:00。",
        "Buffet breakfast is served in the 1st-floor restaurant, 7:00-10:30 on weekdays and 7:00-11:00 on weekends.",
        "周一至五 7:00-10:30；周末 7:00-11:00",
        "Mon-Fri 7:00-10:30; weekends 7:00-11:00",
        "1楼餐厅",
        "1st-floor restaurant",
        "",
        "",
        "86 (21) 1234 5678",
        "儿童1.2米以下免费。（此行为示例，请删除）",
        "Children under 1.2 m eat free. (Example row - please delete)",
    ],
]

# 脱敏示例 xlsx：与 examples/脱敏示例.json 同源的演示数据（非真实酒店）。
# 每语言 3 条回答为同一组事实的三种自然改写（供 TTS 轮换），互补事实写入"其他备注"。
DEMO_ROWS = [
    [
        "1",
        "玩",
        "健身房",
        "Fitness Center",
        "健身房;健身中心;器械",
        "gym;fitness center",
        "健身房在示例酒店2楼，每天6:00到23:00对住店宾客免费开放。\n示例酒店2楼设有健身房，住店宾客可免费使用，开放时间为每天6:00至23:00。\n住店宾客可以免费使用2楼的健身房，每天从早上6:00一直开放到晚上23:00。",
        "The gym is on the 2nd floor of the demo hotel, open daily from 6:00 to 23:00, free for in-house guests.\nLocated on the 2nd floor, the demo hotel's gym is free for in-house guests and open every day from 6:00 to 23:00.\nIn-house guests can use the 2nd-floor gym free of charge, from 6:00 in the morning until 23:00 daily.",
        "每天 6:00-23:00",
        "Daily 6:00-23:00",
        "2楼",
        "2nd floor",
        "乘大堂电梯到2楼即到",
        "Take the lobby elevator to the 2nd floor",
        "86 (21) 1234 5678",
        "配有跑步机、椭圆机等心肺器材和自由重量区，值班教练可提供指导。脱敏演示数据，非真实酒店信息。",
        "Equipped with treadmills, elliptical machines and a free-weights area; a duty trainer can assist. De-identified demo data, not a real hotel.",
    ],
    [
        "2",
        "食",
        "早餐",
        "Breakfast",
        "早餐;自助早餐;几点吃早饭",
        "breakfast;breakfast time",
        "自助早餐在1楼悦餐厅供应，周一至周五7:00-10:30，周末7:00-11:00。\n1楼悦餐厅供应自助早餐，开放时间是周一到周五7:00至10:30、周六和周日7:00至11:00。\n每天的自助早餐都可以在1楼悦餐厅吃到：周一至五7:00开到10:30，周末则从7:00供应到11:00。",
        "Buffet breakfast is served in the 1st-floor Yue Restaurant, 7:00-10:30 on weekdays and 7:00-11:00 on weekends.\nThe 1st-floor Yue Restaurant serves the buffet breakfast, from 7:00 to 10:30 Monday to Friday and 7:00 to 11:00 on weekends.\nYou can have the buffet breakfast at Yue Restaurant on the 1st floor — 7:00-10:30 on weekdays and 7:00-11:00 at weekends.",
        "周一至五 7:00-10:30；周末 7:00-11:00",
        "Mon-Fri 7:00-10:30; weekends 7:00-11:00",
        "1楼悦餐厅",
        "1st-floor restaurant",
        "",
        "",
        "86 (21) 1234 5678",
        "早餐含中式和西式档口，儿童1.2米以下免费；无需预约，凭房卡入场。脱敏演示数据，非真实酒店信息。",
        "Chinese and Western stations; children under 1.2 m eat free; no reservation needed — just bring your room card. De-identified demo data, not a real hotel.",
    ],
    [
        "3",
        "服务",
        "延迟退房",
        "",
        "延迟退房;晚点退房;几点退房",
        "",
        "标准退房时间为中午12:00；如需延迟退房，请在退房当天上午11:00前联系前台，视房态尽量安排，延迟至18:00前通常加收半天房费，具体以前台当日确认为准。\n如需延迟退房，最迟要在退房当天上午11:00前联系前台申请，酒店会视房态尽量安排；标准退房时间是中午12:00，延迟到18:00之前一般加收半天房费，最终以前台当日答复为准。\n退房标准时间为中午12:00，想晚点退房的宾客请在当天上午11:00前向前台提出，酒店将视房态安排；延迟至18:00前通常需加付半天房费，确切口径以当日前台确认为依据。",
        "",
        "退房 12:00；延迟申请截止 11:00",
        "",
        "前台（1楼大堂）",
        "",
        "",
        "",
        "86 (21) 1234 5678",
        "脱敏演示数据；本条故意只填中文，演示缺失语言时英文留空并报告。",
        "",
    ],
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
EXAMPLE_FILL = PatternFill("solid", fgColor="F2F2F2")
WRAP = Alignment(wrap_text=True, vertical="top")


def _style_entry_sheet(sheet, example_rows) -> None:
    sheet.append(ENTRY_COLUMNS)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    for row in example_rows:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            cell.alignment = WRAP
            cell.fill = EXAMPLE_FILL
    widths = [8, 8, 16, 16, 22, 22, 36, 36, 20, 20, 16, 16, 22, 22, 18, 28, 28]
    for index, width in enumerate(widths, start=1):
        column = sheet.cell(row=1, column=index).column_letter
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 42
    sheet.freeze_panes = "A2"

    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(CATEGORIES) + '"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="分类不正确",
        error="请从九个标准分类中选择：衣、食、住、行、游、购、玩、服务、其他",
    )
    sheet.add_data_validation(validation)
    validation.add(f"B2:B{2 + 498}")


def build_workbook(example_rows, extra_note: str = "") -> Workbook:
    workbook = Workbook()

    info = workbook.active
    info.title = "填写说明"
    info.append(["条目", "必填", "说明"])
    for cell in info[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in INSTRUCTIONS:
        info.append(list(row))
    for row in info.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 18
    info.column_dimensions["C"].width = 88
    info.freeze_panes = "A2"

    if extra_note:
        info.append(["本文件说明", "", extra_note])
        for cell in info[info.max_row]:
            cell.alignment = WRAP

    entries = workbook.create_sheet("知识条目")
    _style_entry_sheet(entries, example_rows)

    categories = workbook.create_sheet("分类表")
    categories.append(["分类", "English hint", "包含什么（举例）"])
    for cell in categories[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in CATEGORY_TABLE:
        categories.append(list(row))
    categories.column_dimensions["A"].width = 10
    categories.column_dimensions["B"].width = 22
    categories.column_dimensions["C"].width = 60
    categories.freeze_panes = "A2"
    return workbook


def main() -> int:
    TEMPLATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    EXAMPLE_PATH.parent.mkdir(parents=True, exist_ok=True)

    template = build_workbook(
        TEMPLATE_EXAMPLE_ROWS,
        "灰色两行为示例，正式填写前请删除；一行一个主题，不要合并单元格。",
    )
    template.save(TEMPLATE_PATH)

    demo = build_workbook(
        DEMO_ROWS,
        "脱敏示例（非真实酒店数据），演示与 examples/脱敏示例.json 相同的内容。",
    )
    demo.save(EXAMPLE_PATH)

    print(f"wrote {TEMPLATE_PATH}")
    print(f"wrote {EXAMPLE_PATH}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
